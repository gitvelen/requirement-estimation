import io
import os
import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT_DIR = os.path.dirname(os.path.dirname(__file__))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from backend.service.document_parser import DocumentParser
from backend.utils.embedded_attachment_extractor import EmbeddedAttachmentExtractor


def _build_docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _build_xlsx_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _inject_docx_embedding(host_bytes: bytes, embedded_name: str, embedded_bytes: bytes) -> bytes:
    source = io.BytesIO(host_bytes)
    target = io.BytesIO()
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(f"word/embeddings/{embedded_name}", embedded_bytes)
    return target.getvalue()


def test_document_parser_extracts_embedded_docx_text_from_plain_embedding():
    parser = DocumentParser()
    host_bytes = _build_docx_bytes("主文档正文")
    embedded_bytes = _build_docx_bytes("附件正文一", "附件正文二")
    payload = _inject_docx_embedding(host_bytes, "attachment.docx", embedded_bytes)

    parsed = parser.parse(payload, filename="host.docx", file_type="docx")

    attachment_texts = [item["text"] for item in parsed.get("attachments") or []]
    assert any("附件正文一" in text for text in attachment_texts)
    assert any("附件正文二" in text for text in attachment_texts)


def test_document_parser_extracts_embedded_xlsx_text_from_plain_embedding():
    parser = DocumentParser()
    host_bytes = _build_docx_bytes("主文档正文")
    embedded_bytes = _build_xlsx_bytes([["字段", "说明"], ["支付渠道", "微信支付"]])
    payload = _inject_docx_embedding(host_bytes, "attachment.xlsx", embedded_bytes)

    parsed = parser.parse(payload, filename="host.docx", file_type="docx")

    attachment_texts = [item["text"] for item in parsed.get("attachments") or []]
    assert any("支付渠道" in text and "微信支付" in text for text in attachment_texts)


def test_document_parser_extracts_embedded_txt_text_from_plain_embedding():
    parser = DocumentParser()
    host_bytes = _build_docx_bytes("主文档正文")
    embedded_bytes = "文本附件正文\n第二行补充说明".encode("utf-8")
    payload = _inject_docx_embedding(host_bytes, "attachment.txt", embedded_bytes)

    parsed = parser.parse(payload, filename="host.docx", file_type="docx")

    attachments = parsed.get("attachments") or []
    assert any(item.get("type") == "txt" and "文本附件正文" in item.get("text", "") for item in attachments)


def test_document_parser_limits_recursive_attachment_depth():
    parser = DocumentParser()
    level3 = _build_docx_bytes("第三层正文")
    level2 = _inject_docx_embedding(_build_docx_bytes("第二层正文"), "level3.docx", level3)
    level1 = _inject_docx_embedding(_build_docx_bytes("第一层正文"), "level2.docx", level2)
    host = _inject_docx_embedding(_build_docx_bytes("主文档正文"), "level1.docx", level1)

    parsed = parser.parse(host, filename="host.docx", file_type="docx")

    attachment_texts = [item["text"] for item in parsed.get("attachments") or []]
    merged_text = "\n".join(attachment_texts)
    assert "第一层正文" in merged_text
    assert "第二层正文" in merged_text
    assert "第三层正文" in merged_text


def test_document_parser_isolates_invalid_attachment_failures():
    parser = DocumentParser()
    host_bytes = _build_docx_bytes("主文档正文")
    valid_embedded = _build_docx_bytes("有效附件正文")
    payload = _inject_docx_embedding(host_bytes, "valid.docx", valid_embedded)
    payload = _inject_docx_embedding(payload, "broken.docx", b"not-a-real-docx")

    parsed = parser.parse(payload, filename="host.docx", file_type="docx")

    attachment_texts = [item["text"] for item in parsed.get("attachments") or []]
    errors = parsed.get("attachment_errors") or []
    assert any("有效附件正文" in text for text in attachment_texts)
    assert any("broken.docx" in error.get("name", "") for error in errors)


def test_document_parser_enforces_attachment_count_limit():
    parser = DocumentParser()
    parser._attachment_extractor = EmbeddedAttachmentExtractor(parser._parse_with_context, max_attachment_count=2)

    payload = _build_docx_bytes("主文档正文")
    for idx in range(3):
        payload = _inject_docx_embedding(payload, f"attachment-{idx}.docx", _build_docx_bytes(f"附件正文{idx}"))

    parsed = parser.parse(payload, filename="host.docx", file_type="docx")

    assert len(parsed.get("attachments") or []) == 2
    assert any(error.get("reason") == "attachment-count-limit-exceeded" for error in (parsed.get("attachment_errors") or []))


def test_document_parser_enforces_attachment_size_limit():
    parser = DocumentParser()
    parser._attachment_extractor = EmbeddedAttachmentExtractor(parser._parse_with_context, max_attachment_size=128)

    host_bytes = _build_docx_bytes("主文档正文")
    embedded_bytes = _build_docx_bytes("超大附件正文" * 50)
    payload = _inject_docx_embedding(host_bytes, "too-large.docx", embedded_bytes)

    parsed = parser.parse(payload, filename="host.docx", file_type="docx")

    assert not (parsed.get("attachments") or [])
    assert any("attachment-too-large" in error.get("reason", "") for error in (parsed.get("attachment_errors") or []))


def test_document_parser_extracts_from_real_ole_embedding_sample():
    parser = DocumentParser()
    sample_path = Path("uploads/fca9c523-b4d9-4b7a-aec6-29af800fb762_2026_33_1.docx")

    parsed = parser.parse(sample_path.read_bytes(), filename=sample_path.name, file_type="docx")

    attachments = parsed.get("attachments") or []
    assert attachments, "expected real sample to expose at least one parsed attachment"
    names = [str(item.get("name") or "") for item in attachments]
    assert any("oleObject2" in name for name in names), "expected Workbook stream fallback to recover oleObject2"
    assert any("oleObject5" in name for name in names), "expected WordDocument stream fallback to recover oleObject5"
    errors = parsed.get("attachment_errors") or []
    assert not any("oleObject5.bin" in error.get("name", "") for error in errors), "expected oleObject5 to be recovered instead of degraded"
